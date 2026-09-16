from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "DOCS" / "WIFI_GUARD_개발_진행_및_전체_파이프라인_보고서_20260914.md"
OUTPUT = ROOT / "DOCS" / "WIFI_GUARD_개발_진행_및_전체_파이프라인_보고서_20260914.docx"

FONT = "맑은 고딕"
MONO_FONT = "Consolas"
NAVY = "243B53"
PALE_BLUE = "F3F7FA"
LIGHT_GRAY = "D9D9D9"
CODE_GRAY = "F5F5F5"
BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(255, 255, 255)


def set_run_font(run, name: str = FONT, size: float | None = None, *, bold: bool | None = None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    return run


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def remove_paragraph_borders(paragraph_or_style) -> None:
    p_pr = paragraph_or_style._element.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def set_cell_margins(cell, top: int = 95, start: int = 110, bottom: int = 95, end: int = 110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "5")
        tag.set(qn("w:color"), LIGHT_GRAY)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cannot_split = OxmlElement("w:cantSplit")
    tr_pr.append(cannot_split)


def add_page_number(paragraph) -> None:
    paragraph.add_run("   ")
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, text, end))
    set_run_font(run, size=8)


INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|`.+?`)")


def add_inline(paragraph, text: str, *, size: float | None = None, color=BLACK) -> None:
    cursor = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > cursor:
            run = set_run_font(paragraph.add_run(text[cursor:match.start()]), size=size)
            run.font.color.rgb = color
        token = match.group(0)
        if token.startswith("**"):
            run = set_run_font(paragraph.add_run(token[2:-2]), size=size, bold=True)
        else:
            run = set_run_font(paragraph.add_run(token[1:-1]), MONO_FONT, size=(size or 10.5) - 0.5)
        run.font.color.rgb = color
        cursor = match.end()
    if cursor < len(text):
        run = set_run_font(paragraph.add_run(text[cursor:]), size=size)
        run.font.color.rgb = color


def parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        rows.append([cell.strip() for cell in line.strip().strip("|").split("|")])
    return [rows[0], *rows[2:]]


def table_widths(column_count: int) -> list[float]:
    presets = {
        2: [2.05, 4.95],
        3: [1.65, 2.15, 3.20],
        4: [0.65, 1.55, 2.75, 2.05],
    }
    return presets.get(column_count, [7.0 / column_count] * column_count)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = table_widths(len(rows[0]))
    set_table_borders(table)
    repeat_header(table.rows[0])

    for row_index, (word_row, values) in enumerate(zip(table.rows, rows)):
        prevent_row_split(word_row)
        for column_index, (cell, value) in enumerate(zip(word_row.cells, values)):
            cell.width = Inches(widths[column_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            if row_index == 0:
                set_cell_shading(cell, NAVY)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(paragraph, value, size=9.0, color=WHITE)
                for run in paragraph.runs:
                    run.bold = True
            else:
                if row_index % 2 == 0:
                    set_cell_shading(cell, PALE_BLUE)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if column_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
                add_inline(paragraph, value, size=9.1)

    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def add_code_block(doc: Document, lines: list[str]) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.right_indent = Inches(0.18)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.line_spacing = 1.0
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), CODE_GRAY)
    paragraph._p.get_or_add_pPr().append(shading)
    run = paragraph.add_run("\n".join(lines))
    set_run_font(run, MONO_FONT, 8.3)
    run.font.color.rgb = BLACK


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.70)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.footer_distance = Inches(0.32)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.7)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.22

    title = doc.styles["Title"]
    title.font.name = FONT
    title._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    title.font.size = Pt(26)
    title.font.bold = True
    title.font.color.rgb = BLACK
    title.paragraph_format.space_after = Pt(16)
    remove_paragraph_borders(title)

    for style_name, size, before, after in (
        ("Heading 1", 17, 17, 8),
        ("Heading 2", 13, 13, 6),
    ):
        style = doc.styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(footer.add_run("WIFI GUARD 개발 진행 결과"), size=8)
    add_page_number(footer)
    for run in footer.runs:
        run.font.color.rgb = RGBColor(89, 89, 89)

    doc.core_properties.title = "WIFI GUARD 개발 진행 결과 및 전체 파이프라인 보고서"
    doc.core_properties.subject = "ESP32 C5부터 Kafka Backend Frontend 낙상 모델까지의 개발 현황"
    doc.core_properties.author = "WIFI GUARD 프로젝트"
    doc.core_properties.keywords = "WIFI GUARD CSI ESP32 Raspberry Pi MQTT Kafka"


def build() -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_document(doc)

    index = 0
    cover_metadata = 0
    body_started = False

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            index += 1
            continue

        if stripped.startswith("# "):
            paragraph = doc.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            remove_paragraph_borders(paragraph)
            add_inline(paragraph, stripped[2:], size=26)
            index += 1
            continue

        if stripped.startswith("## "):
            heading_text = stripped[3:].strip()
            if not body_started:
                doc.add_page_break()
                body_started = True
            paragraph = doc.add_paragraph(style="Heading 1")
            add_inline(paragraph, heading_text, size=17)
            index += 1
            continue

        if stripped.startswith("### "):
            paragraph = doc.add_paragraph(style="Heading 2")
            add_inline(paragraph, stripped[4:].strip(), size=13)
            index += 1
            continue

        if stripped.startswith("```"):
            code_lines: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            index += 1
            add_code_block(doc, code_lines)
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and re.match(r"^\|?[\s:|-]+\|", lines[index + 1]):
            table_lines = [line]
            index += 1
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_table(doc, parse_table(table_lines))
            continue

        list_match = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if list_match:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.28)
            paragraph.paragraph_format.first_line_indent = Inches(-0.28)
            paragraph.paragraph_format.space_after = Pt(4)
            if "**" in list_match.group(2):
                paragraph.paragraph_format.keep_with_next = True
            add_inline(paragraph, f"{list_match.group(1)}.  {list_match.group(2)}", size=10.7)
            index += 1
            continue

        if cover_metadata < 3 and not stripped.startswith("##"):
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(7)
            add_inline(paragraph, stripped.rstrip("  "), size=11.5)
            cover_metadata += 1
            index += 1
            continue

        paragraph_lines = [stripped.rstrip("  ")]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if not candidate or candidate.startswith(("#", "|", "```")) or re.match(r"^\d+\.\s+", candidate):
                break
            paragraph_lines.append(candidate.rstrip("  "))
            index += 1
        paragraph = doc.add_paragraph()
        add_inline(paragraph, " ".join(paragraph_lines), size=10.7)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
